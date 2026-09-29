"""Docling-based parsing for layout-rich formats: PDF, DOCX, PPTX, HTML, Markdown.

Docling rebuilds the document structure (headings, lists, tables, reading
order), and its HybridChunker cuts on that structure: chunks never straddle a
heading, tables stay whole, and each chunk carries its heading path.
"""
from __future__ import annotations

from functools import lru_cache
from pathlib import Path

from langchain_core.documents import Document

from rag.ingestion.detect import Detection
from rag.ingestion.schema import FileType, SourceInfo, make_doc

MAX_TOKENS = 512


@lru_cache
def _converter():
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    # OCR off: pages without a text layer go to the vision model instead,
    # which reads handwriting far better than Docling's built-in OCR.
    pdf_opts = PdfPipelineOptions(do_ocr=False, do_table_structure=True)
    return DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_opts)})


@lru_cache
def _chunker():
    import tiktoken
    from docling.chunking import HybridChunker
    from docling_core.transforms.chunker.tokenizer.openai import OpenAITokenizer

    # Count tokens with the same tokenizer family the OpenAI embedding model uses.
    tokenizer = OpenAITokenizer(tokenizer=tiktoken.get_encoding("cl100k_base"), max_tokens=MAX_TOKENS)
    return HybridChunker(tokenizer=tokenizer, merge_peers=True)


def _docling_chunks(path: Path, src: SourceInfo, parser: str) -> list[Document]:
    dl_doc = _converter().convert(str(path)).document
    chunker = _chunker()
    docs = []
    for chunk in chunker.chunk(dl_doc):
        text = chunker.contextualize(chunk=chunk)   # prepends the heading path
        if not text.strip():
            continue
        pages = sorted({prov.page_no for item in chunk.meta.doc_items for prov in (item.prov or [])})
        headings = chunk.meta.headings or []
        docs.append(make_doc(text, src, parser=parser, page=pages[0] if pages else None,
                             section_path=" > ".join(headings), prechunked=True,
                             pages=pages or None))
    return docs


def parse_pdf(path: Path, det: Detection, src: SourceInfo, ctx) -> list[Document]:
    if ctx.settings.pdf_parser == "fast":
        docs = _fast_pdf_pages(path, det, src)
    else:
        docs = _docling_chunks(path, src, "pdf_text")
    if det.scanned_pages:
        # Mixed PDF: Docling found nothing on these pages, so read them visually.
        from rag.ingestion.parsers.ocr_vision import ocr_pdf_pages
        docs += ocr_pdf_pages(path, det.scanned_pages, src, ctx)
    return docs


def _fast_pdf_pages(path: Path, det: Detection, src: SourceInfo) -> list[Document]:
    """PDF_PARSER=fast: the text layer page by page, no layout model.

    Several times cheaper per page than Docling. Headings, reading order of
    multi-column pages and table structure are not recovered; the chunk step
    then splits each page by tokens.
    """
    import pypdfium2 as pdfium

    scanned = set(det.scanned_pages)
    docs = []
    pdf = pdfium.PdfDocument(str(path))
    try:
        for n in range(1, len(pdf) + 1):
            if n in scanned:
                continue
            page = pdf[n - 1]
            textpage = page.get_textpage()
            text = textpage.get_text_range().strip()
            textpage.close()
            page.close()
            if text:
                docs.append(make_doc(text, src, parser="pdf_fast", page=n, section_path=f"Page {n}"))
    finally:
        pdf.close()
    return docs


def parse_office(path: Path, det: Detection, src: SourceInfo, ctx) -> list[Document]:
    docs = _docling_chunks(path, src, "office")
    if det.file_type is FileType.pptx:
        docs += _pptx_speaker_notes(path, src)
    return docs


def parse_web(path: Path, det: Detection, src: SourceInfo, ctx) -> list[Document]:
    return _docling_chunks(path, src, "web_markup")


def _pptx_speaker_notes(path: Path, src: SourceInfo) -> list[Document]:
    """Docling reads slide text only; speaker notes often hold the real explanation."""
    from pptx import Presentation

    docs = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        if not slide.has_notes_slide:
            continue
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            title = slide.shapes.title.text if slide.shapes.title is not None else f"Slide {n}"
            docs.append(make_doc(f"{title} (speaker notes)\n{notes}", src, parser="office",
                                 page=n, section_path=f"{title} > Speaker notes"))
    return docs
