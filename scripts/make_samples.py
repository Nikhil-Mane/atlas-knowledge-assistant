"""Generate the binary one-page sample files in data/samples/.

Text formats (md, html, csv, json, js, txt) are committed as-is; this script
builds the formats that need a library: EML, PDF (text), PDF (scanned),
PNG, DOCX, PPTX and XLSX.

Usage:  .venv/Scripts/python scripts/make_samples.py
"""
from __future__ import annotations

import random
from email.message import EmailMessage
from email.utils import format_datetime
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont

OUT = Path(__file__).resolve().parent.parent / "data" / "samples"
FONTS = Path("C:/Windows/Fonts")


def font(name: str, size: int) -> ImageFont.FreeTypeFont:
    try:
        return ImageFont.truetype(str(FONTS / name), size)
    except OSError:
        return ImageFont.load_default(size)


# --------------------------------------------------------------------- EML
def make_email() -> None:
    msg = EmailMessage()
    msg["From"] = "Priya Sharma <priya.sharma@example.com>"
    msg["To"] = "backend-team@example.com"
    msg["Subject"] = "Incident summary: memory leak in order-service"
    msg["Date"] = format_datetime(datetime(2026, 9, 14, 9, 30, tzinfo=timezone.utc))
    msg.set_content(
        "Hi team,\n\n"
        "Short summary of Friday's incident in order-service.\n\n"
        "What happened: memory grew from 300 MB to 1.8 GB over 6 hours and the\n"
        "pods were OOM-killed. Logs showed 'MaxListenersExceededWarning: Possible\n"
        "EventEmitter memory leak detected. 11 message listeners added'.\n\n"
        "Root cause: every HTTP request called redisSubscriber.on('message', ...)\n"
        "and never removed the listener, so listeners and their closures piled up.\n\n"
        "Fix: register the listener once at startup, and use emitter.once() or\n"
        "emitter.off() for per-request listeners. We also added a heap snapshot\n"
        "alert using 'node --heapsnapshot-near-heap-limit=2'.\n\n"
        "The full postmortem is attached.\n\n"
        "Thanks,\nPriya\n"
    )
    msg.add_attachment(
        (
            "POSTMORTEM - order-service memory leak\n"
            "Detection: Grafana alert on container memory > 1.5 GB.\n"
            "Time to detect: 6 hours. Time to mitigate: 40 minutes (rollback).\n"
            "Action items:\n"
            " 1. Add a lint rule that flags .on() inside request handlers.\n"
            " 2. Load test every release for 2 hours and compare heap size.\n"
            " 3. Capture heap snapshots with Chrome DevTools and diff them.\n"
        ).encode(),
        maintype="text",
        subtype="plain",
        filename="postmortem.txt",
    )
    (OUT / "07-incident-email.eml").write_bytes(bytes(msg))


# ---------------------------------------------------------------- PDF text
def make_text_pdf() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(OUT / "08-modules-guide.pdf"), pagesize=A4,
                            title="CommonJS vs ES Modules", author="Sample corpus")
    story = [
        Paragraph("CommonJS vs ES Modules in Node.js", styles["Title"]),
        Paragraph(
            "Node.js supports two module systems. CommonJS (CJS) is the original "
            "system and uses <b>require()</b> and <b>module.exports</b>. ES Modules (ESM) "
            "are the JavaScript standard and use <b>import</b> and <b>export</b>.",
            styles["BodyText"]),
        Spacer(1, 10),
        Paragraph("How Node.js decides which system a file uses", styles["Heading2"]),
        Paragraph(
            "Files ending in <b>.mjs</b> are always ESM and <b>.cjs</b> are always CommonJS. "
            "For <b>.js</b> files, Node.js reads the nearest package.json: "
            "<b>\"type\": \"module\"</b> means ESM, otherwise CommonJS.",
            styles["BodyText"]),
        Spacer(1, 10),
        Paragraph("Key differences", styles["Heading2"]),
    ]
    table = Table([
        ["Feature", "CommonJS", "ES Modules"],
        ["Syntax", "require / module.exports", "import / export"],
        ["Loading", "synchronous", "asynchronous"],
        ["Top-level await", "not supported", "supported"],
        ["__dirname", "available", "use import.meta.dirname"],
        ["File extension in imports", "optional", "required ('./util.js')"],
    ], colWidths=[140, 160, 170])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563C9")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
    ]))
    story += [
        table,
        Spacer(1, 12),
        Paragraph("Loading ESM from CommonJS", styles["Heading2"]),
        Paragraph(
            "Older Node.js versions throw <b>ERR_REQUIRE_ESM</b> when require() loads an ES "
            "module. Use a dynamic import instead, which works in both systems:",
            styles["BodyText"]),
        Preformatted("const { default: chalk } = await import('chalk');", styles["Code"]),
    ]
    doc.build(story)


# ----------------------------------------------------------- PDF scanned
def make_scanned_pdf() -> None:
    """Image-only PDF with handwriting-style text, like a phone scan of notes."""
    random.seed(7)
    w, h = 1240, 1754  # A4 at 150 dpi
    page = Image.new("RGB", (w, h), (246, 243, 232))
    d = ImageDraw.Draw(page)
    for y in range(180, h - 80, 64):  # ruled notebook lines
        d.line([(60, y), (w - 60, y)], fill=(170, 195, 225), width=2)
    d.line([(150, 0), (150, h)], fill=(230, 150, 150), width=2)  # margin

    title = font("Inkfree.ttf", 64)
    hand = font("segoepr.ttf", 38)
    ink = (28, 45, 110)
    d.text((180, 95), "Buffers in Node.js", font=title, fill=ink)
    lines = [
        "- Buffer = fixed-size chunk of raw binary memory",
        "- lives OUTSIDE the V8 heap",
        "- Buffer.alloc(10)  -> 10 bytes, zero-filled (safe)",
        "- Buffer.allocUnsafe(10) -> faster, old data inside!",
        "- Buffer.from('hello', 'utf8')",
        "- buf.toString('base64')  /  'hex'",
        "- 1 char != 1 byte  (utf8 emoji = 4 bytes)",
        "- Buffer.byteLength('héllo') = 6",
        "- new Buffer() is DEPRECATED -> never use",
        "- streams give data as Buffers by default",
        "",
        "Remember: allocUnsafe only if you overwrite",
        "every byte right away!!",
    ]
    y = 205
    for text in lines:
        x = 180 + random.randint(-6, 10)
        d.text((x, y + random.randint(-4, 4)), text, font=hand, fill=ink)
        y += 64

    # Scanner look: slight rotation, blur and grain.
    page = page.rotate(-1.2, expand=False, fillcolor=(236, 233, 222), resample=Image.BICUBIC)
    page = page.filter(ImageFilter.GaussianBlur(0.6))
    noise = Image.effect_noise((w, h), 18).convert("RGB")
    page = Image.blend(page, noise, 0.06)
    page.save(OUT / "09-buffers-handwritten-scan.pdf", "PDF", resolution=150)


# ---------------------------------------------------------------- PNG
def make_whiteboard_png() -> None:
    w, h = 1400, 800
    img = Image.new("RGB", (w, h), (250, 250, 247))
    d = ImageDraw.Draw(img)
    marker = font("segoepr.ttf", 34)
    small = font("segoepr.ttf", 26)
    big = font("Inkfree.ttf", 56)
    blue, red, green = (30, 70, 170), (190, 40, 40), (20, 120, 70)

    d.text((60, 40), "Scaling with the cluster module", font=big, fill=blue)
    # primary
    d.rounded_rectangle([560, 150, 840, 250], 18, outline=red, width=5)
    d.text((600, 175), "Primary", font=marker, fill=red)
    d.text((440, 262), "cluster.fork() x os.availableParallelism()", font=small, fill=red)
    # workers
    for i, x in enumerate([160, 480, 800, 1120]):
        d.rounded_rectangle([x, 400, x + 200, 490], 16, outline=blue, width=4)
        d.text((x + 28, 425), f"Worker {i + 1}", font=marker, fill=blue)
        d.line([(700, 300), (x + 100, 398)], fill=(90, 90, 90), width=3)
    d.text((140, 530), "each worker = separate process + own event loop", font=marker, fill=green)
    d.text((140, 590), "all share port 3000 (primary distributes connections)", font=marker, fill=green)
    d.text((140, 650), "worker crash -> cluster.on('exit') -> fork a new one", font=marker, fill=green)
    d.text((140, 710), "PM2 does this for you:  pm2 start app.js -i max", font=marker, fill=blue)
    img.filter(ImageFilter.GaussianBlur(0.4)).save(OUT / "10-cluster-whiteboard.png")


# ---------------------------------------------------------------- DOCX
def make_docx() -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.title = "Error Handling Guidelines"
    doc.add_heading("Error Handling Guidelines for Node.js Services", level=1)
    doc.add_paragraph(
        "These guidelines apply to every backend service owned by the platform team. "
        "Follow them so that errors are logged once, reported clearly and never crash "
        "a process unexpectedly.")

    doc.add_heading("1. Operational vs programmer errors", level=2)
    doc.add_paragraph(
        "Operational errors are expected runtime problems: a timeout, a refused "
        "connection, invalid user input. Handle them and return a proper response.")
    doc.add_paragraph(
        "Programmer errors are bugs: reading a property of undefined, passing the "
        "wrong argument type. Log them, let the process crash and let the process "
        "manager restart it.")

    doc.add_heading("2. Rules", level=2)
    for rule in [
        "Always await promises or attach .catch(); unhandled rejections crash the process.",
        "Throw Error objects (or subclasses), never strings, so the stack trace is kept.",
        "Add context when re-throwing: new Error('failed to load user', { cause: err }).",
        "Handle errors in one Express error middleware instead of in every route.",
        "Register process.on('uncaughtException') only to log and exit, never to continue.",
    ]:
        doc.add_paragraph(rule, style="List Bullet")

    doc.add_heading("3. HTTP status codes", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Situation"
    table.rows[0].cells[1].text = "Status"
    for situation, status in [
        ("Validation failed", "400 Bad Request"),
        ("Missing or invalid token", "401 Unauthorized"),
        ("Resource does not exist", "404 Not Found"),
        ("Duplicate unique value", "409 Conflict"),
        ("Unexpected bug", "500 Internal Server Error"),
        ("Downstream service timed out", "503 Service Unavailable"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text = situation, status

    for p in doc.paragraphs:
        for r in p.runs:
            r.font.size = Pt(11) if p.style.name.startswith(("Normal", "List")) else r.font.size
    doc.save(OUT / "11-error-handling.docx")


# ---------------------------------------------------------------- PPTX
def make_pptx() -> None:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide.shapes.title.text = "Async Patterns in Node.js"
    body = slide.placeholders[1].text_frame
    points = [
        ("Callbacks", "error-first: (err, result) => {} - leads to nesting"),
        ("Promises", "chain with .then(); util.promisify() converts callbacks"),
        ("async / await", "reads like sync code; wrap in try/catch"),
        ("Promise.all", "run in parallel, fails fast on the first rejection"),
        ("Promise.allSettled", "waits for all, reports each success or failure"),
    ]
    body.text = f"{points[0][0]}: {points[0][1]}"
    for title, text in points[1:]:
        p = body.add_paragraph()
        p.text = f"{title}: {text}"
    for p in body.paragraphs:
        for r in p.runs:
            r.font.size = Pt(20)
    slide.notes_slide.notes_text_frame.text = (
        "Speaker notes: avoid 'await' inside a for loop when the calls are independent - "
        "collect the promises and use Promise.all so they run concurrently. Use "
        "p-limit when you need to cap concurrency, for example 5 requests at a time.")
    prs.save(OUT / "12-async-patterns.pptx")


# ---------------------------------------------------------------- XLSX
def make_xlsx() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "LTS schedule"
    ws.append(["Node.js release schedule (even-numbered releases become LTS)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    header = ["Version", "Codename", "Released", "Active LTS start", "End of life"]
    ws.append(header)
    for row in [
        ["18", "Hydrogen", "2022-04-19", "2022-10-25", "2025-04-30"],
        ["20", "Iron", "2023-04-18", "2023-10-24", "2026-04-30"],
        ["22", "Jod", "2024-04-24", "2024-10-29", "2027-04-30"],
        ["24", "Krypton", "2025-05-06", "2025-10-28", "2028-04-30"],
    ]:
        ws.append(row)
    fill = PatternFill("solid", fgColor="2563C9")
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    for col, width in zip("ABCDE", [10, 12, 14, 18, 14]):
        ws.column_dimensions[col].width = width
    ws.append([])
    ws.append(["Note: odd-numbered releases are never LTS and are supported for about 6 months."])

    notes = wb.create_sheet("Upgrade notes")
    notes.append(["From", "To", "Watch out for"])
    notes["A1"].font = notes["B1"].font = notes["C1"].font = Font(bold=True)
    notes.append(["18", "20", "Stable test runner (node:test); permission model added (experimental)"])
    notes.append(["20", "22", "require(esm) support; built-in WebSocket client; node --run"])
    notes.append(["22", "24", "npm 11; URLPattern global; check native addons rebuild"])
    notes.column_dimensions["C"].width = 70
    wb.save(OUT / "13-lts-schedule.xlsx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for fn in (make_email, make_text_pdf, make_scanned_pdf, make_whiteboard_png,
               make_docx, make_pptx, make_xlsx):
        fn()
        print(f"ok  {fn.__name__}")
